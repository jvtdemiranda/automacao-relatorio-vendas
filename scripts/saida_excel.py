"""
Relatório em Excel, pra quem quer filtrar e conferir os números.

Quatro abas, cada uma se explicando sozinha (título, o que é, como ler):
Resumo · Vendas por dia · Todas as vendas · Qualidade dos dados.
"""

import io
import math
import re
import zipfile
from datetime import datetime

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.table import Table, TableStyleInfo

from analise import brl

MARCA, CINZA, TINTA = "1F4E5A", "5F6368", "1B1F23"
VERDE, VERMELHO = "1E7F4F", "B3261E"
FUNDO = PatternFill("solid", start_color="EEF3F4", end_color="EEF3F4")
FUNDO_MARCA = PatternFill("solid", start_color=MARCA, end_color=MARCA)
LINHA = Side(style="thin", color="D5DDDF")
BRL = '"R$" #,##0.00;[Red]-"R$" #,##0.00'
PCT = "0.0%;[Red]-0.0%"
DATA_FIXA = datetime(2026, 1, 1)
DIAS_SEMANA = ["segunda", "terça", "quarta", "quinta", "sexta", "sábado", "domingo"]


# ---------------------------------------------------------------- utilidades

def texto(ws, linha, coluna, valor):
    """
    Texto que veio das planilhas das lojas é gravado SEMPRE como texto. Sem
    isso, um produto cadastrado como "=HYPERLINK(...)" viraria uma fórmula
    de verdade no Excel do gerente (injeção de fórmula).
    """
    c = ws.cell(row=linha, column=coluna, value=None if valor is None else str(valor))
    c.data_type = "s"
    return c


def larguras(ws, valores):
    for j, w in enumerate(valores, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w


def paragrafo(ws, linha, valor, colunas, fonte=None, caracteres_por_linha=None):
    """Texto longo numa faixa mesclada, com quebra de linha e altura ajustada."""
    ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=colunas)
    c = ws.cell(row=linha, column=1, value=valor)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if fonte:
        c.font = fonte
    largura = caracteres_por_linha or sum(ws.column_dimensions[get_column_letter(j)].width for j in range(1, colunas + 1))
    ws.row_dimensions[linha].height = 15.5 * max(1, math.ceil(len(valor) * 1.08 / largura))
    return linha + 1


def titulo(ws, nome, subtitulo, colunas):
    ws.cell(row=1, column=1, value=nome).font = Font(size=17, bold=True, color=MARCA)
    ws.row_dimensions[1].height = 26
    paragrafo(ws, 2, subtitulo, colunas, Font(italic=True, color=CINZA))
    return 4


def secao(ws, linha, nome, colunas):
    for j in range(1, colunas + 1):
        ws.cell(row=linha, column=j).border = Border(bottom=Side(style="medium", color=MARCA))
    ws.cell(row=linha, column=1, value=nome).font = Font(size=12.5, bold=True, color=MARCA)
    ws.row_dimensions[linha].height = 20
    return linha + 1


def cabecalho(ws, linha, nomes):
    for j, nome in enumerate(nomes, start=1):
        c = ws.cell(row=linha, column=j, value=nome)
        c.fill = FUNDO_MARCA
        c.font = Font(bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[linha].height = 32


def linha_tabela(ws, linha, n_colunas, zebra, negrito=False):
    for j in range(1, n_colunas + 1):
        c = ws.cell(row=linha, column=j)
        c.border = Border(bottom=LINHA)
        c.alignment = Alignment(horizontal=c.alignment.horizontal, vertical="top", wrap_text=c.alignment.wrap_text)
        if zebra:
            c.fill = FUNDO
        if negrito:
            c.font = Font(bold=True)


def para_impressao(ws, paisagem=True):
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if paisagem else "portrait"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def grafico_barras(titulo_grafico, eixo_y_formato="#,##0"):
    g = BarChart()
    g.title = titulo_grafico
    g.style = 10
    g.y_axis.numFmt = eixo_y_formato
    g.y_axis.majorGridlines = None
    # O openpyxl 3.1 não grava que os eixos são visíveis, e as versões
    # recentes do Excel escondem eixo sem essa marcação: o gráfico sairia sem
    # os nomes das lojas e sem os valores.
    g.x_axis.delete = False
    g.y_axis.delete = False
    g.height, g.width = 7.5, 15
    return g


# ---------------------------------------------------------------- abas

def aba_resumo(wb, r):
    ws = wb.active
    ws.title = "Resumo"
    ws.sheet_view.showGridLines = False
    larguras(ws, [24, 17, 17, 11, 13, 9, 14, 14, 27])
    n = 9
    linha = titulo(ws, f"Relatório de vendas · {r['mes_extenso'].capitalize()}",
                   "Farmácia Vitalis Manipulação · consolidado das lojas Centro, Umarizal e Ananindeua · "
                   "gerado automaticamente a partir das planilhas enviadas pelas lojas", n)

    linha = secao(ws, linha, "Destaques do mês", n)
    for frase in r["destaques"]:
        linha = paragrafo(ws, linha, "• " + frase, n,
                          Font(bold=frase.startswith("Atenção"), color=VERMELHO if frase.startswith("Atenção") else TINTA))
    linha += 1

    # Indicadores principais: 4 blocos, como no PDF
    t = r["total"]
    blocos = [
        ("FATURAMENTO LÍQUIDO", t["liquido"], BRL, None),
        (f"VS {r['mes_anterior_extenso'].upper()}", t["variacao"], PCT, 0),
        ("DA META CONSOLIDADA", t["pct_meta"], PCT, 1),
        ("TICKET MÉDIO", t["ticket_medio"], BRL, None),
    ]
    for i, (rotulo, valor, fmt, limite) in enumerate(blocos):
        c1, c2 = [(1, 2), (3, 4), (5, 6), (7, 9)][i]
        for lin in (linha, linha + 1):
            ws.merge_cells(start_row=lin, start_column=c1, end_row=lin, end_column=c2)
            for j in range(c1, c2 + 1):
                ws.cell(row=lin, column=j).fill = FUNDO
        ws.cell(row=linha, column=c1, value=rotulo).font = Font(size=8.5, bold=True, color=CINZA)
        v = ws.cell(row=linha + 1, column=c1, value=valor if valor is not None else "—")
        v.number_format = fmt
        cor = TINTA if limite is None or valor is None else (VERDE if valor >= limite else VERMELHO)
        v.font = Font(size=16, bold=True, color=cor)
        v.alignment = Alignment(horizontal="left")
    ws.row_dimensions[linha + 1].height = 26
    linha += 3

    linha = secao(ws, linha, "Resultado por loja", n)
    mes_ant = r["mes_anterior_extenso"].split("/")[0]
    cabecalho(ws, linha, ["Loja", "Faturamento líquido", "Meta", "% da meta", f"Variação vs {mes_ant}",
                          "Vendas", "Ticket médio", "Devoluções", "Situação"])
    inicio_lojas = linha + 1
    for i, l in enumerate(r["lojas"] + [dict(r["total"], loja="Total", enviou=True)]):
        linha += 1
        total = l["loja"] == "Total"
        ws.cell(row=linha, column=1, value=l["loja"])
        for col, chave, fmt in [(2, "liquido", BRL), (3, "meta", BRL), (4, "pct_meta", PCT), (5, "variacao", PCT),
                                (6, "vendas", "0"), (7, "ticket_medio", BRL), (8, "devolucoes", BRL)]:
            ws.cell(row=linha, column=col, value=l.get(chave)).number_format = fmt
        if not l["enviou"]:
            situacao, cor = "Sem dados — ver Qualidade", VERMELHO
        elif l["pct_meta"] is None:
            situacao, cor = "Sem meta cadastrada", CINZA
        elif l["pct_meta"] >= 1:
            situacao, cor = "Meta batida", VERDE
        else:
            situacao, cor = f"Faltaram {brl(l['meta'] - l['liquido'])}", VERMELHO
        linha_tabela(ws, linha, n, zebra=i % 2 == 1 and not total, negrito=total)
        s = ws.cell(row=linha, column=9, value=situacao)
        s.font = Font(bold=True, color=cor)
        s.alignment = Alignment(indent=1, vertical="top")
        if total:
            for j in range(1, n + 1):
                ws.cell(row=linha, column=j).border = Border(top=Side(style="thin", color=CINZA))
    fim_lojas = linha
    for coluna, limite in [("D", 1), ("E", 0)]:
        faixa = f"{coluna}{inicio_lojas}:{coluna}{fim_lojas}"
        ws.conditional_formatting.add(faixa, CellIsRule(operator="lessThan", formula=[str(limite)], font=Font(color=VERMELHO, bold=True)))
        ws.conditional_formatting.add(faixa, CellIsRule(operator="greaterThanOrEqual", formula=[str(limite)], font=Font(color=VERDE, bold=True)))
    linha += 2

    linha = secao(ws, linha, "Faturamento por categoria", n)
    cabecalho(ws, linha, ["Categoria", "Faturamento líquido", "Participação"])
    inicio_cat = linha + 1
    for i, c in enumerate(r["categorias"]):
        linha += 1
        ws.cell(row=linha, column=1, value=c["categoria"])
        ws.cell(row=linha, column=2, value=c["liquido"]).number_format = BRL
        ws.cell(row=linha, column=3, value=c["participacao"]).number_format = PCT
        linha_tabela(ws, linha, 3, zebra=i % 2 == 1)
    fim_cat = linha
    linha += 2

    # Gráficos (mesmos do PDF), lado a lado — começando numa página nova na
    # impressão, pra não serem cortados ao meio
    ws.row_breaks.append(Break(id=linha - 1))
    g1 = grafico_barras("Faturamento líquido x meta por loja")
    g1.add_data(Reference(ws, min_col=2, max_col=3, min_row=inicio_lojas - 1, max_row=fim_lojas - 1), titles_from_data=True)
    g1.set_categories(Reference(ws, min_col=1, min_row=inicio_lojas, max_row=fim_lojas - 1))
    g1.series[0].graphicalProperties.solidFill = MARCA
    g1.series[1].graphicalProperties.solidFill = "C9D3D6"
    ws.add_chart(g1, f"A{linha}")

    g2 = grafico_barras("Faturamento por categoria")
    g2.type = "bar"
    g2.legend = None
    g2.add_data(Reference(ws, min_col=2, min_row=inicio_cat - 1, max_row=fim_cat), titles_from_data=True)
    g2.set_categories(Reference(ws, min_col=1, min_row=inicio_cat, max_row=fim_cat))
    g2.series[0].graphicalProperties.solidFill = MARCA
    g2.x_axis.scaling.orientation = "maxMin"  # maior categoria em cima, como na tabela
    ws.add_chart(g2, f"E{linha}")
    linha += 16

    linha = secao(ws, linha, "Como ler estes números", n)
    for termo, explicacao in [
        ("Faturamento líquido", "tudo o que foi vendido no mês, menos as devoluções."),
        ("% da meta", "faturamento líquido dividido pela meta da loja; 100% ou mais = meta batida."),
        (f"Variação vs {mes_ant}", "quanto o faturamento líquido subiu ou caiu em relação ao mês anterior."),
        ("Vendas", "número de atendimentos (cupons ou pedidos) com pelo menos um item vendido."),
        ("Ticket médio", "valor médio de cada venda: total vendido ÷ número de vendas."),
        ("Devoluções", "valor devolvido pelos clientes no mês, já descontado do faturamento líquido."),
        ("De onde vêm os dados", "das planilhas enviadas pelas lojas; a aba \"Qualidade dos dados\" mostra o "
                                 "que precisou ser corrigido em cada uma."),
    ]:
        ws.cell(row=linha, column=1, value=termo).font = Font(bold=True)
        ws.merge_cells(start_row=linha, start_column=2, end_row=linha, end_column=n)
        c = ws.cell(row=linha, column=2, value=explicacao)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=linha, column=1).alignment = Alignment(vertical="top")
        linha += 1
    para_impressao(ws)


def aba_diario(wb, r, dados):
    ws = wb.create_sheet("Vendas por dia")
    ws.sheet_view.showGridLines = False
    lojas = [l["loja"] for l in r["lojas"]]
    colunas = ["Data", "Dia da semana"] + lojas + ["Total do dia"]
    larguras(ws, [13, 15] + [15] * len(lojas) + [16])
    n = len(colunas)
    linha = titulo(ws, f"Vendas por dia · {r['mes_extenso'].capitalize()}",
                   "Faturamento líquido de cada loja em cada dia com movimento (vendas menos devoluções). "
                   "Domingo as lojas não abrem.", n)
    cabecalho(ws, linha, colunas)
    ws.freeze_panes = ws.cell(row=linha + 1, column=1)
    ws.print_title_rows = f"{linha}:{linha}"
    inicio = linha + 1

    por_dia = dados.pivot_table(index="data", columns="loja", values="valor", aggfunc="sum", fill_value=0.0)
    for i, (dia, valores) in enumerate(por_dia.sort_index().iterrows()):
        linha += 1
        ws.cell(row=linha, column=1, value=datetime(dia.year, dia.month, dia.day)).number_format = "dd/mm/yyyy"
        ws.cell(row=linha, column=2, value=DIAS_SEMANA[dia.weekday()])
        for j, loja in enumerate(lojas, start=3):
            ws.cell(row=linha, column=j, value=round(float(valores.get(loja, 0.0)), 2)).number_format = BRL
        ws.cell(row=linha, column=n, value=f"=SUM(C{linha}:{get_column_letter(n - 1)}{linha})").number_format = BRL
        ws.cell(row=linha, column=n).font = Font(bold=True)
        linha_tabela(ws, linha, n, zebra=i % 2 == 1)
    fim = linha

    linha += 1
    ws.cell(row=linha, column=1, value="Total do mês")
    for j in range(3, n + 1):
        letra = get_column_letter(j)
        ws.cell(row=linha, column=j, value=f"=SUM({letra}{inicio}:{letra}{fim})").number_format = BRL
    linha_tabela(ws, linha, n, zebra=False, negrito=True)
    for j in range(1, n + 1):
        ws.cell(row=linha, column=j).border = Border(top=Side(style="thin", color=CINZA))

    if fim >= inicio:
        g = grafico_barras("Faturamento líquido por dia, por loja")
        g.grouping = "stacked"
        g.overlap = 100
        g.add_data(Reference(ws, min_col=3, max_col=n - 1, min_row=inicio - 1, max_row=fim), titles_from_data=True)
        g.set_categories(Reference(ws, min_col=1, min_row=inicio, max_row=fim))
        g.x_axis.number_format = "dd"
        g.x_axis.title = "dia do mês"
        g.height, g.width = 9, 22
        for serie, cor in zip(g.series, [MARCA, "5B9AA8", "C9D3D6"]):
            serie.graphicalProperties.solidFill = cor
        ws.add_chart(g, f"{get_column_letter(n + 2)}{inicio - 1}")
    para_impressao(ws)


def aba_itens(wb, r, dados):
    ws = wb.create_sheet("Todas as vendas")
    colunas = ["Data", "Loja", "Nº da venda", "Tipo", "Categoria", "Produto", "Qtd", "Valor"]
    larguras(ws, [13, 13, 13, 11, 24, 38, 6, 13])
    n = len(colunas)
    linha = titulo(ws, f"Todas as vendas · {r['mes_extenso'].capitalize()}",
                   "Cada linha é um item vendido (ou devolvido), já corrigido e no mesmo formato para as três "
                   "lojas. Use as setas do cabeçalho para filtrar por loja, categoria ou produto. "
                   "Devoluções aparecem com valor negativo, em vermelho.", n)
    cabecalho(ws, linha, colunas)
    ws.print_title_rows = f"{linha}:{linha}"
    topo = linha
    ordenado = dados.sort_values(["data", "loja", "venda"], kind="stable")
    for l in ordenado.itertuples(index=False):
        linha += 1
        ws.cell(row=linha, column=1, value=datetime(l.data.year, l.data.month, l.data.day)).number_format = "dd/mm/yyyy"
        texto(ws, linha, 2, l.loja)
        texto(ws, linha, 3, l.venda)
        ws.cell(row=linha, column=4, value="Devolução" if l.valor < 0 else "Venda")
        texto(ws, linha, 5, l.categoria)
        texto(ws, linha, 6, l.produto)
        ws.cell(row=linha, column=7, value=int(l.qtd))
        ws.cell(row=linha, column=8, value=float(l.valor)).number_format = BRL
    if len(ordenado):
        tabela = Table(displayName="vendas", ref=f"A{topo}:{get_column_letter(n)}{linha}")
        tabela.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(tabela)
    ws.freeze_panes = ws.cell(row=topo + 1, column=1)
    para_impressao(ws)


# O que cada tipo de problema significa, em linguagem de quem lê o relatório:
# (o que foi encontrado, o que o robô fez, por que importa)
DESCARTES = {
    "linhas de total": ("A planilha já vinha com uma linha de TOTAL no fim", "Linha ignorada na soma",
                        lambda v: f"Se fosse somada, a loja apareceria com {brl(v)} a mais (o dobro)" if v
                        else "Se fosse somada, o faturamento da loja sairia dobrado"),
    "linhas duplicadas": ("A mesma venda aparecia duas vezes no arquivo", "Cópias removidas",
                          lambda v: f"Evitou somar {brl(v)} a mais" if v else "Evitou somar vendas em dobro"),
    "ilegível": ("Data, valor ou quantidade impossível de ler", "Linhas deixadas de fora",
                 lambda v: "Conferir essas vendas com a loja"),
    "fora do mês": ("Vendas de outro mês dentro do arquivo", "Deixadas de fora deste relatório",
                    lambda v: f"{brl(v)} que pertencem a outro mês" if v else "Pertencem a outro mês"),
    "número errado de colunas": ("Linha com um \";\" a mais no texto, desalinhando as colunas", "Linhas deixadas de fora",
                                 lambda v: "Conferir essas vendas com a loja"),
}


def ocorrencias(q):
    """Lista (o que foi encontrado, linhas, o que o robô fez, por que importa) de uma loja."""
    itens = []
    for motivo, n in q["descartes"].items():
        chave = next((k for k in DESCARTES if k in motivo), None)
        if chave:
            encontrado, acao, impacto = DESCARTES[chave]
            itens.append((encontrado, n, acao, impacto(q["valor_descartado"].get(motivo, 0.0))))
        else:
            itens.append((motivo, n, "Deixadas de fora", ""))
    if q["devolucoes"]:
        itens.append(("Devoluções lançadas com valor negativo", q["devolucoes"], "Descontadas do faturamento",
                      f"Faturamento líquido já sem os {brl(q['valor_devolucoes'])} devolvidos"))
    for aviso in q["avisos"]:
        if "codificação Windows" in aviso:
            itens.append(("Arquivo salvo com a acentuação antiga do Windows", None, "Convertido automaticamente",
                          "Sem isso, os acentos viriam trocados e as categorias não seriam reconhecidas"))
        elif "categoria não reconhecida" in aviso:
            n = int(aviso.rsplit(":", 1)[-1])
            itens.append(("Categoria com um nome que o robô não conhece", n, "Somadas em \"Outros\"",
                          "Cadastrar o nome novo no robô para entrar na categoria certa"))
        elif "NÃO ENVIADA" in aviso:
            itens.append(("Planilha do mês não recebida", None, "—", "O total do mês NÃO inclui esta loja"))
        elif "arquivos encontrados" in aviso:
            itens.append(("Mais de um arquivo desta loja no mês", None, "Nenhum foi usado, para não somar em dobro",
                          "O total do mês NÃO inclui esta loja — deixar só o arquivo certo na pasta"))
        else:
            itens.append((aviso[0].upper() + aviso[1:], None, "Loja fora do total", "O total do mês NÃO inclui esta loja"))
    return itens


def aba_qualidade(wb, r):
    ws = wb.create_sheet("Qualidade dos dados")
    ws.sheet_view.showGridLines = False
    larguras(ws, [14, 40, 12, 30, 52])
    n = 5
    linha = titulo(ws, "Qualidade dos dados recebidos",
                   "Cada loja manda a planilha num formato diferente. Antes de somar, o robô conferiu cada arquivo, "
                   "corrigiu o que dava para corrigir e deixou de fora o que não dava. Esta aba mostra tudo o que "
                   "foi feito — para ninguém usar um número sem saber de onde ele veio.", n)

    linha = secao(ws, linha, "Resumo por loja", n)
    cabecalho(ws, linha, ["Loja", "Arquivo recebido", "Linhas no arquivo", "Usadas no relatório", "Situação"])
    for i, q in enumerate(r["qualidade"]):
        linha += 1
        itens = ocorrencias(q)
        fora_do_total = not q["arquivo"] or not q["linhas_validas"]
        if fora_do_total:
            situacao, cor = "Atenção: loja FORA do total do mês", VERMELHO
        elif itens:
            situacao, cor = f"Corrigida automaticamente ({len(itens)} {'ajuste' if len(itens) == 1 else 'ajustes'})", TINTA
        else:
            situacao, cor = "Sem problemas", VERDE
        texto(ws, linha, 1, q["loja"]).font = Font(bold=True)
        texto(ws, linha, 2, q["arquivo"] or "nenhum arquivo recebido")
        ws.cell(row=linha, column=3, value=q["linhas_lidas"] or "—").alignment = Alignment(horizontal="center")
        usadas = ws.cell(row=linha, column=4,
                         value=f"{q['linhas_validas']} de {q['linhas_lidas']}" if q["linhas_lidas"] else "—")
        usadas.alignment = Alignment(horizontal="center")
        ws.cell(row=linha, column=5, value=situacao).font = Font(bold=True, color=cor)
        linha_tabela(ws, linha, n, zebra=i % 2 == 1)
    linha += 2

    linha = secao(ws, linha, "O que foi encontrado em cada planilha", n)
    cabecalho(ws, linha, ["Loja", "O que foi encontrado", "Linhas", "O que o robô fez", "Por que importa"])
    i = 0
    for q in r["qualidade"]:
        itens = ocorrencias(q) or [("Nenhum problema encontrado", None, "—", "Planilha usada como veio")]
        for encontrado, qtd, acao, impacto in itens:
            linha += 1
            texto(ws, linha, 1, q["loja"]).font = Font(bold=True)
            for col, valor in [(2, encontrado), (4, acao), (5, impacto)]:
                c = texto(ws, linha, col, valor)
                c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(row=linha, column=3, value=qtd if qtd is not None else "—").alignment = Alignment(horizontal="center", vertical="top")
            maior = max(len(encontrado) / 40, len(acao) / 28, len(impacto) / 50)
            ws.row_dimensions[linha].height = 15.5 * max(1, math.ceil(maior))
            linha_tabela(ws, linha, n, zebra=i % 2 == 1)
            i += 1
    linha += 2
    paragrafo(ws, linha, "\"Linhas no arquivo\" conta as linhas de venda do arquivo, sem o cabeçalho e os títulos. "
                         "Nomes de clientes que vêm em algumas planilhas são descartados na leitura (LGPD) e não "
                         "aparecem em nenhuma parte deste relatório.", n, Font(italic=True, size=9, color=CINZA))
    para_impressao(ws)


# ---------------------------------------------------------------- arquivo

def salvar(caminho, resumo, dados):
    wb = Workbook()
    aba_resumo(wb, resumo)
    aba_diario(wb, resumo, dados)
    aba_itens(wb, resumo, dados)
    aba_qualidade(wb, resumo)
    wb.active = 0
    wb.properties.creator = "automacao-relatorio-vendas"
    wb.properties.title = f"Relatório de vendas {resumo['mes_extenso']}"
    wb.properties.created = wb.properties.modified = DATA_FIXA

    # O .xlsx é um zip, e o openpyxl carimba a hora atual em cada arquivo
    # interno e na data de "modificado" (ignorando a que definimos acima).
    # Regravar com data fixa deixa o arquivo idêntico a cada execução: assim
    # o robô que publica os relatórios só commita quando algo mudou.
    buffer = io.BytesIO()
    wb.save(buffer)
    with zipfile.ZipFile(buffer) as origem, zipfile.ZipFile(caminho, "w", zipfile.ZIP_DEFLATED) as destino:
        for item in origem.infolist():
            conteudo = origem.read(item.filename)
            if item.filename == "docProps/core.xml":
                conteudo = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>2026-01-01T00:00:00Z", conteudo)
            info = zipfile.ZipInfo(item.filename, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            destino.writestr(info, conteudo)
