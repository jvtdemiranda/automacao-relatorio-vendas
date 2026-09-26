"""
Gera as planilhas de exemplo em entrada/AAAA-MM/, simulando o que as três
lojas de uma farmácia de manipulação (fictícia) mandam todo fim de mês —
cada uma exportada de um sistema diferente, com os problemas reais que
isso traz:

- Centro (Excel): três linhas de título antes da tabela, coluna com o nome
  do cliente e uma linha "TOTAL GERAL" no fim.
- Umarizal (CSV): codificação do Windows (cp1252), separador ";", vírgula
  decimal, data dd/mm/aaaa e devoluções com valor negativo.
- Ananindeua (Excel): colunas com outros nomes, valor como texto
  ("R$ 1.234,56"), categorias escritas de outro jeito e linhas duplicadas
  (a mesma venda exportada duas vezes).

Período e semente fixos: rodar de novo gera exatamente os mesmos dados.
"""

import csv
import os
import random
from datetime import date, timedelta

from openpyxl import Workbook

RAIZ = os.path.join(os.path.dirname(__file__), "..")
MESES = [(2026, 7), (2026, 8)]

PRODUTOS = {
    "Dermatológicos": [
        ("Minoxidil 5% loção 60ml", 68, 95), ("Ácido hialurônico sérum 30ml", 85, 140),
        ("Vitamina C 15% sérum 30ml", 79, 120), ("Creme clareador 30g", 90, 180),
        ("Protetor solar FPS 50 com cor 50g", 75, 110),
    ],
    "Cápsulas e suplementos": [
        ("Vitamina D3 10.000UI 60 cáps", 45, 70), ("Ômega 3 1000mg 60 cáps", 60, 95),
        ("Colágeno + vitamina C 60 cáps", 70, 115), ("Magnésio dimalato 60 cáps", 55, 85),
        ("Complexo B 60 cáps", 40, 65),
    ],
    "Florais e homeopatia": [
        ("Floral de Bach 30ml", 28, 45), ("Arnica CH6 glóbulos", 25, 40), ("Fórmula floral personalizada", 35, 60),
    ],
    "Veterinários": [
        ("Suplemento articular pet 30 cáps", 55, 90), ("Pasta palatável pet 30 doses", 60, 110),
    ],
    "Revenda": [
        ("Hidratante corporal 200ml", 25, 55), ("Sabonete facial 150ml", 18, 40), ("Protetor labial", 12, 25),
    ],
}

# Peso de cada categoria por loja (mix diferente em cada bairro)
MIX = {
    "centro": [30, 35, 12, 5, 18],
    "umarizal": [42, 30, 10, 3, 15],
    "ananindeua": [18, 38, 14, 12, 18],
}
VENDAS_POR_DIA = {"centro": 16, "umarizal": 13, "ananindeua": 10}
CRESCIMENTO_AGOSTO = {"centro": 1.06, "umarizal": 0.95, "ananindeua": 1.12}

# Como cada sistema escreve as categorias
CATEGORIA_ANANINDEUA = {
    "Dermatológicos": "DERMATO", "Cápsulas e suplementos": "Suplemento",
    "Florais e homeopatia": "florais/homeo", "Veterinários": "VET", "Revenda": "revenda ",
}

NOMES = ["Ana", "Bruno", "Carla", "Diego", "Elaine", "Fábio", "Gabriela", "Hugo", "Isabela", "João",
         "Karina", "Lucas", "Márcia", "Nelson", "Olívia", "Paulo", "Rita", "Sérgio", "Tânia", "Vítor"]


def dias_uteis(ano, mes):
    d = date(ano, mes, 1)
    while d.month == mes:
        if d.weekday() < 6:  # farmácia fecha domingo
            yield d
        d += timedelta(days=1)


def vendas_da_loja(loja, ano, mes, primeiro_numero):
    """Lista de itens vendidos: (data, nº da venda, categoria, produto, qtd, valor)."""
    categorias = list(PRODUTOS)
    fator = CRESCIMENTO_AGOSTO[loja] if mes == 8 else 1.0
    itens, numero = [], primeiro_numero
    for dia in dias_uteis(ano, mes):
        base = VENDAS_POR_DIA[loja] * fator * (0.6 if dia.weekday() == 5 else 1.0)
        for _ in range(max(1, round(random.gauss(base, base * 0.18)))):
            numero += 1
            for _ in range(random.choices([1, 2, 3], weights=[65, 27, 8])[0]):
                categoria = random.choices(categorias, weights=MIX[loja])[0]
                produto, minimo, maximo = random.choice(PRODUTOS[categoria])
                qtd = random.choices([1, 2], weights=[88, 12])[0]
                valor = round(random.uniform(minimo, maximo), 2) * qtd
                itens.append((dia, numero, categoria, produto, qtd, round(valor, 2)))
    return itens, numero


def br(valor):
    return f"{valor:.2f}".replace(".", ",")


def salvar_centro(pasta, ano, mes, itens):
    wb = Workbook()
    ws = wb.active
    ws.title = "Vendas"
    ws.append(["Farmácia Vitalis Manipulação - Loja Centro"])
    ws.append([f"Relatório de vendas - {mes:02d}/{ano}"])
    ws.append([])
    ws.append(["Data", "Nº Venda", "Cliente", "Categoria", "Descrição", "Qtd", "Valor Total"])
    clientes = {}
    for dia, numero, categoria, produto, qtd, valor in itens:
        cliente = clientes.setdefault(numero, f"{random.choice(NOMES)} {random.choice('ABCDEFGHLMNPRST')}.")
        ws.append([dia, numero, cliente, categoria, produto, qtd, valor])
        ws.cell(row=ws.max_row, column=1).number_format = "dd/mm/yyyy"
    ws.append(["TOTAL GERAL", None, None, None, None, sum(i[4] for i in itens), round(sum(i[5] for i in itens), 2)])
    wb.save(os.path.join(pasta, f"vendas_centro_{ano}-{mes:02d}.xlsx"))


def salvar_umarizal(pasta, ano, mes, itens):
    linhas = []
    for dia, numero, categoria, produto, qtd, valor in itens:
        linhas.append([dia.strftime("%d/%m/%Y"), numero, categoria.upper(), produto, qtd, br(valor)])
        if random.random() < 0.015:  # devolução registrada depois, com valor negativo
            dia_dev = min(dia + timedelta(days=random.randint(1, 5)), date(ano, mes, 28))
            if dia_dev.weekday() == 6:  # domingo a loja está fechada
                dia_dev += timedelta(days=1)
            linhas.append([dia_dev.strftime("%d/%m/%Y"), numero, categoria.upper(),
                           f"DEVOLUÇÃO - {produto}", -qtd, br(-valor)])
    with open(os.path.join(pasta, f"umarizal_vendas_{ano}{mes:02d}.csv"), "w", newline="", encoding="cp1252") as f:
        escritor = csv.writer(f, delimiter=";")
        escritor.writerow(["DATA", "CUPOM", "CATEGORIA", "PRODUTO", "QUANTIDADE", "VALOR"])
        escritor.writerows(linhas)


def salvar_ananindeua(pasta, ano, mes, itens):
    wb = Workbook()
    ws = wb.active
    ws.append(["dt_venda", "num_pedido", "tipo_produto", "item", "qtde", "total (R$)"])
    for dia, numero, categoria, produto, qtd, valor in itens:
        linha = [dia.strftime("%d/%m/%Y"), f"AN-{numero}", CATEGORIA_ANANINDEUA[categoria], produto, qtd,
                 "R$ " + f"{valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")]
        ws.append(linha)
        if random.random() < 0.03:  # exportação sobreposta: mesma linha duas vezes
            ws.append(linha)
    nome_mes = ["JAN", "FEV", "MAR", "ABR", "MAI", "JUN", "JUL", "AGO", "SET", "OUT", "NOV", "DEZ"][mes - 1]
    wb.save(os.path.join(pasta, f"Vendas Ananindeua {nome_mes}-{str(ano)[2:]}.xlsx"))


def main():
    random.seed(77)
    numero = {"centro": 31000, "umarizal": 58000, "ananindeua": 12000}
    for ano, mes in MESES:
        pasta = os.path.join(RAIZ, "entrada", f"{ano}-{mes:02d}")
        os.makedirs(pasta, exist_ok=True)
        for loja, salvar in [("centro", salvar_centro), ("umarizal", salvar_umarizal), ("ananindeua", salvar_ananindeua)]:
            itens, numero[loja] = vendas_da_loja(loja, ano, mes, numero[loja])
            salvar(pasta, ano, mes, itens)
            print(f"{ano}-{mes:02d} {loja}: {len(itens)} itens, R$ {sum(i[5] for i in itens):,.2f}")

    with open(os.path.join(RAIZ, "entrada", "metas.csv"), "w", newline="", encoding="utf-8") as f:
        escritor = csv.writer(f, delimiter=";")
        escritor.writerow(["mes", "loja", "meta"])
        for mes, metas in [("2026-07", (45000, 38000, 26000)), ("2026-08", (47000, 40000, 27000))]:
            for loja, meta in zip(["Centro", "Umarizal", "Ananindeua"], metas):
                escritor.writerow([mes, loja, meta])


if __name__ == "__main__":
    main()
